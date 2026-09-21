"""Regenerate the dashboard snapshot .xlsx from the scored pool + applied overrides.

Run before every update to data/mc_overrides_applied.json or
data/gm_overrides_applied.json, per the market-cap refresh workflow (see
docs/handoff-brief.md). Uses scripts/scoring.py for the v6 model so this
never drifts from what scripts/build_site.py publishes.

Writes a dated snapshot to snapshots/shortlist_snapshot_<date>.xlsx, plus a
compact shortlist_snapshot_<date>.json (same data, lean field names) that
scripts/build_site.py embeds into the dashboard for the "view an older
date" history feature.
"""
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from scoring import ROOT, compact_snapshot_rows, load_companies, load_overrides, score_pool

SNAPSHOTS = ROOT / "snapshots"


def build_rows():
    companies = load_companies()
    mc_overrides, gm_overrides = load_overrides()
    scored, _ = score_pool(companies, mc_overrides, gm_overrides)

    rows = []
    for c, s in zip(companies, scored):
        rows.append({
            'ticker': c['ticker'], 'name': c['name'], 'country': c['country'],
            'mc': round(s['mc'], 2), 'mc_is_updated': c['ticker'] in mc_overrides,
            'equity': c['equity'], 'pb': round(s['pb'], 3) if s['pb'] else None,
            'roe': round(s['roe'], 4) if s['roe'] is not None else None,
            'rev_growth': round(c['revGrowth'], 4) if c['revGrowth'] is not None else None,
            'gm': s['gm'], 'gm_is_fallback': s['gm_is_fallback'],
            'gm_fiscal_year': c['gmFiscalYear'], 'gm_source_type': c['gmSourceType'],
            'rev_growth_pts': c['fixedPts']['revenue_growth'],
            'rev_consistency_pts': c['fixedPts']['revenue_consistency'],
            'dilution_pts': c['fixedPts']['dilution'],
            'valuation_pts': round(s['val_pts'], 2),
            'profitability_pts': c['fixedPts']['profitability'],
            'cfo_pts': c['fixedPts']['cfo_positive'],
            'track_record_pts': c['fixedPts']['track_record'],
            'dividend_pts': c['fixedPts']['dividend_consistency'],
            'gm_pts': round(s['gm_pts'], 2),
            'total_score': s['total'],
        })

    rows.sort(key=lambda r: -r['total_score'])
    for i, r in enumerate(rows):
        r['rank'] = i + 1
    return rows


def write_workbook(rows, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Snapshot'

    snap_dt = datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')
    ws['A1'] = 'European Small/Mid-Cap Value Shortlist — Dashboard Snapshot'
    ws['A1'].font = Font(name='Arial', size=14, bold=True)
    ws['A2'] = f'Captured {snap_dt}, before batch market-cap/GM update'
    ws['A2'].font = Font(name='Arial', size=10, italic=True, color='555555')
    ws['A3'] = f'{len(rows)} companies in candidate pool (percentile-ranked against full pool)'
    ws['A3'].font = Font(name='Arial', size=10, color='555555')

    headers = ['Rank', 'Ticker', 'Company', 'Country', 'Mkt Cap ($M)', 'MC updated?',
               'P/B', 'ROE', 'Rev Growth', 'Gross Margin', 'GM fallback?', 'GM FY', 'GM Source',
               'RevGrowth pts', 'RevConsist pts', 'Dilution pts', 'Valuation pts',
               'Profitability pts', 'CFO+ pts', 'TrackRecord pts', 'Dividend pts', 'GM pts',
               'v6 Total Score']
    header_row = 5
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=j, value=h)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2F4B3C')
        cell.alignment = Alignment(horizontal='center')

    field_order = ['rank', 'ticker', 'name', 'country', 'mc', 'mc_is_updated', 'pb', 'roe', 'rev_growth',
                   'gm', 'gm_is_fallback', 'gm_fiscal_year', 'gm_source_type',
                   'rev_growth_pts', 'rev_consistency_pts', 'dilution_pts', 'valuation_pts',
                   'profitability_pts', 'cfo_pts', 'track_record_pts', 'dividend_pts', 'gm_pts', 'total_score']
    for i, r in enumerate(rows, start=header_row + 1):
        for j, f in enumerate(field_order, start=1):
            ws.cell(row=i, column=j, value=r[f])

    widths = [6, 14, 32, 16, 12, 11, 7, 8, 10, 11, 11, 7, 16, 12, 12, 11, 12, 14, 9, 13, 11, 7, 12]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + j) if j <= 26 else 'A' + chr(64 + j - 26)].width = w

    ws.freeze_panes = f'A{header_row + 1}'
    wb.save(out_path)


def write_history_json(out_path):
    """Compact per-company records for the dashboard's history dropdown,
    built straight from scoring.compact_snapshot_rows so this can never
    drift from what build_site.py embeds for the same date.
    """
    companies = load_companies()
    mc_overrides, gm_overrides = load_overrides()
    compact = compact_snapshot_rows(companies, mc_overrides, gm_overrides)
    with open(out_path, 'w') as f:
        json.dump(compact, f, separators=(',', ':'))


if __name__ == '__main__':
    rows = build_rows()
    date_str = sys.argv[1] if len(sys.argv) > 1 else datetime.now(timezone.utc).strftime('%Y-%m-%d')
    SNAPSHOTS.mkdir(exist_ok=True)
    xlsx_path = SNAPSHOTS / f'shortlist_snapshot_{date_str}.xlsx'
    json_path = SNAPSHOTS / f'shortlist_snapshot_{date_str}.json'
    write_workbook(rows, xlsx_path)
    write_history_json(json_path)
    print(f'saved {len(rows)} rows -> {xlsx_path} and {json_path}')
